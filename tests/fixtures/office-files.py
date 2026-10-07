"""Synthetic round-trip checks, run inside a fresh offline sandbox as uid 1000."""
import hashlib
import json
import os
from pathlib import Path
import subprocess

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
import pdfplumber
from reportlab.pdfgen import canvas

root = Path("/home/agent/workspace")
inputs = root / "uploads" / "synthetic"
outputs = root / "office-output" / "synthetic"
inputs.mkdir(parents=True)
outputs.mkdir(parents=True)
assert os.getuid() == 1000
subprocess.run(["/opt/portal/office-check"], check=True)
image = inputs / "image.png"
Image.new("RGB", (16, 16), color="blue").save(image)

doc = Document()
doc.add_heading("Office fixture", 0)
doc.add_paragraph("Original paragraph")
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Name"
table.cell(1, 0).text = "Alice"
doc.add_picture(str(image), width=Inches(1))
word = inputs / "source.docx"
doc.save(word)
doc = Document(word)
doc.paragraphs[1].runs[0].text = "Edited paragraph"
edited_word = outputs / "edited.docx"
doc.save(edited_word)
check = Document(edited_word)
assert check.paragraphs[1].text == "Edited paragraph"
assert check.tables[0].cell(1, 0).text == "Alice"
assert len(check.inline_shapes) == 1

book = Workbook()
sheet = book.active
sheet.title = "Budget"
for row in [("Item", "Amount"), ("First", 10), ("Second", 20)]:
    sheet.append(row)
sheet["A4"] = "Total"
sheet["B4"] = "=SUM(B2:B3)"
table = Table(displayName="BudgetData", ref="A1:B3")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
sheet.add_table(table)
chart = BarChart()
chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
sheet.add_chart(chart, "D2")
excel = inputs / "budget.xlsx"
book.save(excel)
book = load_workbook(excel)
book["Budget"]["B2"] = 15
edited_excel = outputs / "edited-budget.xlsx"
book.save(edited_excel)
assert load_workbook(edited_excel)["Budget"]["B4"].value == "=SUM(B2:B3)"
assert len(load_workbook(edited_excel)["Budget"]._charts) == 1

slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text = "Office presentation"
slide.placeholders[1].text = "Original slide"
slide.shapes.add_picture(str(image), Inches(6), Inches(1), width=Inches(1))
table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(3), Inches(1)).table
table.cell(0, 0).text = "Summary"
data = CategoryChartData()
data.categories = ["First", "Second"]
data.add_series("Amount", (15, 20))
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(3), Inches(4), Inches(3), data)
powerpoint = inputs / "slides.pptx"
slides.save(powerpoint)
slides = Presentation(powerpoint)
slides.slides[0].placeholders[1].text = "Edited slide"
edited_slides = outputs / "edited-slides.pptx"
slides.save(edited_slides)
check = Presentation(edited_slides)
assert check.slides[0].placeholders[1].text == "Edited slide"
assert any(shape.has_table for shape in check.slides[0].shapes)
assert any(shape.has_chart for shape in check.slides[0].shapes)

pdf = inputs / "source.pdf"
c = canvas.Canvas(str(pdf))
c.drawString(50, 750, "Office PDF fixture")
for x in [50, 150, 250]:
    c.line(x, 600, x, 660)
for y in [600, 630, 660]:
    c.line(50, y, 250, y)
for x, y, text in [(60, 640, "Name"), (160, 640, "Amount"), (60, 610, "Alice"), (160, 610, "35")]:
    c.drawString(x, y, text)
c.showPage()
c.drawString(50, 750, "Second page")
c.save()
with pdfplumber.open(pdf) as reader:
    assert reader.pages[0].extract_tables()[0] == [["Name", "Amount"], ["Alice", "35"]]
reader = PdfReader(pdf)
writer = PdfWriter()
for page in reader.pages:
    writer.add_page(page)
writer.add_page(reader.pages[0])
merged = outputs / "merged.pdf"
writer.write(merged)
assert len(PdfReader(merged).pages) == 3
writer = PdfWriter()
writer.add_page(reader.pages[1])
split = outputs / "split.pdf"
writer.write(split)
assert "Second page" in PdfReader(split).pages[0].extract_text()
csv = inputs / "data.csv"
csv.write_text("Name,Amount\nAlice,35\n", encoding="utf-8")
hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs.iterdir()}


def convert(source, target_format, folder, **options):
    return subprocess.run(["/opt/portal/office-convert", str(source), "--to", target_format,
                           "--outdir", str(folder)], capture_output=True, text=True, timeout=100, **options)


recalculated = outputs / "recalculated"
result = convert(edited_excel, "xlsx", recalculated)
assert result.returncode == 0, result.stdout + result.stderr
assert load_workbook(recalculated / "edited-budget.xlsx", data_only=True)["Budget"]["B4"].value == 35
csv_result = convert(csv, "xlsx", outputs / "csv")
assert csv_result.returncode == 0, csv_result.stdout + csv_result.stderr
assert load_workbook(outputs / "csv" / "data.xlsx").active["A2"].value == "Alice"
for source in [edited_word, edited_excel, edited_slides]:
    folder = outputs / (source.stem + "-pdf")
    result = convert(source, "pdf", folder)
    assert result.returncode == 0, result.stdout + result.stderr
    exported = PdfReader(folder / (source.stem + ".pdf"))
    assert len(exported.pages) >= 1
    assert any(page.extract_text().strip() for page in exported.pages)

# Reject overwrites, inputs outside the workspace, corrupt documents, and uploads output folders.
assert convert(edited_excel, "xlsx", outputs).returncode != 0
assert convert(edited_word, "pdf", inputs).returncode != 0
assert convert("/etc/passwd", "pdf", outputs / "bad").returncode != 0
corrupt = inputs / "corrupt.docx"
corrupt.write_bytes(b"not a document")
assert convert(corrupt, "pdf", outputs / "corrupt").returncode != 0

# A fake conversion process gives a deterministic timeout check without depending on document size.
fake = outputs / "fake-bin"
fake.mkdir()
(fake / "libreoffice").write_text("#!/bin/sh\n/bin/sleep 3\n", encoding="utf-8")
(fake / "libreoffice").chmod(0o755)
env = {**os.environ, "PATH": str(fake)}
timed = subprocess.run(["/opt/portal/office-convert", str(edited_word), "--to", "pdf", "--outdir",
                        str(outputs / "timeout"), "--timeout", "1"], env=env, capture_output=True, text=True)
assert timed.returncode == 1 and "timed out" in timed.stdout
assert all(hashlib.sha256(path.read_bytes()).hexdigest() == digest for path, digest in hashes.items())
print(json.dumps({"ok": True, "word": True, "excel": True, "powerpoint": True, "pdf": True, "originals_preserved": True}))
